"""Extração conservadora de linhas de disciplinas em históricos PDF e DOCX."""

from io import BytesIO
import re
import unicodedata
from pathlib import PurePath

from docx import Document
from pypdf import PdfReader


CAMPOS = ("codigo", "nome", "nota", "carga_horaria", "periodo", "situacao")
_CODIGO = re.compile(r"^(?P<codigo>(?=[A-Z0-9._/-]*[A-Z])[A-Z0-9][A-Z0-9._/-]{2,})\b\s*(?:[-–:]\s*)?(?P<resto>.*)$", re.I)
_PERIODO = re.compile(r"(?<!\d)(?:(?:19|20)\d{2}\s*/\s*[12]|[12]\s*/\s*(?:19|20)\d{2})(?!\d)")
_NOTA_ROTULADA = re.compile(r"\b(?:nota|m[eé]dia|conceito)\s*[:=-]?\s*(\d{1,2}(?:[,.]\d{1,2})?)\b", re.I)
_CARGA_ROTULADA = re.compile(r"\b(?:carga\s*hor[aá]ria|ch)\s*[:=-]?\s*(\d{1,3})\s*(?:h|hs|horas?)?\b", re.I)
_CARGA_COM_UNIDADE = re.compile(r"\b(\d{1,3})\s*(?:h|hs|horas?)\b", re.I)
_SITUACAO = re.compile(
    r"\b(aprovad[oa]|reprovad[oa]|cursando|em\s+curso|dispensad[oa]|"
    r"aproveitad[oa]|trancad[oa]|cancelad[oa]|reprovad[oa]\s+por\s+nota)\b", re.I
)


def _sem_acentos(valor):
    valor = unicodedata.normalize("NFKD", valor or "")
    return "".join(c for c in valor if not unicodedata.combining(c)).lower().strip()


def _normalizar_campo(celula):
    texto = _sem_acentos(celula)
    if "cod" in texto:
        return "codigo"
    if any(palavra in texto for palavra in ("disciplina", "componente", "nome")):
        return "nome"
    if any(palavra in texto for palavra in ("nota", "media", "conceito")):
        return "nota"
    if any(palavra in texto for palavra in ("carga horaria", "horas", "ch")):
        return "carga_horaria"
    if any(palavra in texto for palavra in ("periodo", "semestre", "ano letivo", "ano/semestre")):
        return "periodo"
    if any(palavra in texto for palavra in ("situacao", "status", "resultado")):
        return "situacao"
    return None


def _separar_colunas(linha):
    if "\t" in linha:
        return [parte.strip() for parte in linha.split("\t")]
    if "|" in linha:
        return [parte.strip() for parte in linha.strip("| ").split("|")]
    return [parte.strip() for parte in re.split(r"\s{2,}", linha.strip()) if parte.strip()]


def _criar_registro(campos, evidencia):
    registro = {campo: (campos.get(campo) or "").strip() or None for campo in CAMPOS}
    registro["texto_origem"] = evidencia.strip()
    if registro["codigo"] and not registro["nome"]:
        correspondencia = _CODIGO.match(registro["codigo"])
        if correspondencia:
            registro["codigo"] = correspondencia.group("codigo")
            registro["nome"] = correspondencia.group("resto").strip() or None
    if not registro["nome"] and not registro["codigo"]:
        return None
    return registro


def _ler_linha_com_colunas(colunas, cabecalho, evidencia):
    if not cabecalho or len(colunas) != len(cabecalho):
        return None
    campos = {}
    for indice, campo in enumerate(cabecalho):
        if campo and colunas[indice]:
            campos[campo] = colunas[indice]
    if campos.get("nome"):
        nome = campos["nome"]
        codigo = _CODIGO.match(nome)
        if codigo:
            campos.setdefault("codigo", codigo.group("codigo"))
            campos["nome"] = codigo.group("resto").strip() or None
    return _criar_registro(campos, evidencia)


def _ler_linha_sem_cabecalho(linha):
    texto = " ".join(linha.strip().split())
    if not texto:
        return None
    codigo_match = _CODIGO.match(texto)
    if not codigo_match:
        return None
    codigo = codigo_match.group("codigo")
    resto = codigo_match.group("resto")
    campos = {"codigo": codigo}

    periodo = _PERIODO.search(resto)
    if periodo:
        campos["periodo"] = periodo.group(0).replace(" ", "")
        resto = resto[:periodo.start()] + " " + resto[periodo.end():]

    nota = _NOTA_ROTULADA.search(resto)
    if nota:
        campos["nota"] = nota.group(1)
        resto = resto[:nota.start()] + " " + resto[nota.end():]

    carga = _CARGA_ROTULADA.search(resto) or _CARGA_COM_UNIDADE.search(resto)
    if carga:
        campos["carga_horaria"] = carga.group(1)
        resto = resto[:carga.start()] + " " + resto[carga.end():]

    situacao = _SITUACAO.search(resto)
    if situacao:
        campos["situacao"] = situacao.group(1)
        resto = resto[:situacao.start()] + " " + resto[situacao.end():]

    nome = re.sub(r"\s+", " ", resto).strip(" -–|\t")
    if nome:
        campos["nome"] = nome
    return _criar_registro(campos, texto)


def _interpretar_linhas(linhas):
    registros = []
    cabecalho = None
    for linha in linhas:
        if not linha or not linha.strip():
            continue
        colunas = _separar_colunas(linha)
        mapa = [_normalizar_campo(coluna) for coluna in colunas]
        if "nome" in mapa and sum(campo is not None for campo in mapa) >= 2:
            cabecalho = mapa
            continue
        registro = _ler_linha_com_colunas(colunas, cabecalho, linha)
        if registro is None:
            registro = _ler_linha_sem_cabecalho(linha)
        if registro:
            registros.append(registro)
    return registros


def extrair_disciplinas(nome_arquivo, conteudo):
    extensao = PurePath(nome_arquivo or "").suffix.lower()
    try:
        if extensao == ".pdf":
            pdf = PdfReader(BytesIO(conteudo), strict=False)
            linhas = []
            for pagina in pdf.pages:
                texto = pagina.extract_text(extraction_mode="layout") or ""
                linhas.extend(texto.splitlines())
        elif extensao == ".docx":
            documento = Document(BytesIO(conteudo))
            linhas = [paragrafo.text for paragrafo in documento.paragraphs]
            for tabela in documento.tables:
                for linha in tabela.rows:
                    linhas.append("\t".join(celula.text.strip() for celula in linha.cells))
        else:
            raise ValueError("Formato não suportado. Envie um arquivo PDF ou DOCX.")
    except ValueError:
        raise
    except Exception as erro:
        raise ValueError("Não foi possível ler o documento. Verifique se o PDF/DOCX não está corrompido ou protegido.") from erro

    resultados = _interpretar_linhas(linhas)
    aviso = None
    if not any(linha.strip() for linha in linhas):
        aviso = "Não foi possível extrair texto. O PDF pode ser uma imagem digitalizada; os dados não foram inferidos."
    elif not resultados:
        aviso = "O texto foi lido, mas nenhuma linha de disciplina pôde ser identificada com segurança."
    return resultados, aviso
